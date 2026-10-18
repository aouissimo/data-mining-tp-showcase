from __future__ import annotations

import re
from dataclasses import dataclass
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.enum.text import MSO_AUTO_SIZE, PP_ALIGN
from pptx.util import Inches, Pt


SOURCE_MD = Path("presentation.md")
OUTPUT_PPTX = Path("presentation.pptx")


SLIDE_BG = RGBColor(8, 18, 36)
SLIDE_BG_ALT = RGBColor(15, 23, 42)
TEXT_MAIN = RGBColor(248, 250, 252)
TEXT_MUTED = RGBColor(191, 219, 254)
TEXT_BODY = RGBColor(226, 232, 240)
ACCENT = RGBColor(96, 165, 250)
ACCENT_SOFT = RGBColor(30, 64, 175)
PANEL_FILL = RGBColor(17, 24, 39)


@dataclass
class SlideContent:
    title: str
    subtitle: str
    bullets: list[tuple[int, str]]
    extras: list[str]


def strip_markdown(text: str) -> str:
    text = re.sub(r"`([^`]+)`", r"\1", text)
    text = re.sub(r"\*\*([^*]+)\*\*", r"\1", text)
    text = re.sub(r"\*([^*]+)\*", r"\1", text)
    return text.strip()


def parse_presentation(markdown_path: Path) -> list[SlideContent]:
    content = markdown_path.read_text(encoding="utf-8")
    content = re.sub(r"<!--.*?-->", "", content, flags=re.S)
    raw_slides = [
        section.strip()
        for section in re.split(r"^\s*---\s*$", content, flags=re.M)
        if section.strip()
    ]

    slides: list[SlideContent] = []
    for raw_slide in raw_slides:
        title = ""
        subtitle = ""
        bullets: list[tuple[int, str]] = []
        extras: list[str] = []

        for line in raw_slide.splitlines():
            stripped = line.rstrip()
            if not stripped:
                continue

            if stripped.startswith("# "):
                title = strip_markdown(stripped[2:])
                continue

            if stripped.startswith("## ") or stripped.startswith("### "):
                header_text = strip_markdown(stripped.split(" ", 1)[1])
                if not subtitle:
                    subtitle = header_text
                else:
                    extras.append(header_text)
                continue

            bullet_match = re.match(r"^(\s*)\*\s+(.*)$", stripped)
            if bullet_match:
                indent = len(bullet_match.group(1))
                level = min(indent // 2, 2)
                bullets.append((level, strip_markdown(bullet_match.group(2))))
                continue

            extras.append(strip_markdown(stripped))

        slides.append(
            SlideContent(
                title=title,
                subtitle=subtitle,
                bullets=bullets,
                extras=extras,
            )
        )

    return slides


def add_background(slide) -> None:
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = SLIDE_BG

    top_band = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(0.58)
    )
    top_band.fill.solid()
    top_band.fill.fore_color.rgb = SLIDE_BG_ALT
    top_band.line.fill.background()

    bottom_band = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(0), Inches(7.03), Inches(13.333), Inches(0.47)
    )
    bottom_band.fill.solid()
    bottom_band.fill.fore_color.rgb = ACCENT_SOFT
    bottom_band.line.fill.background()


def add_textbox(slide, left: float, top: float, width: float, height: float, text: str, *, font_name: str, font_size: int, color: RGBColor, bold: bool = False, align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = box.text_frame
    tf.word_wrap = True
    tf.text = text
    paragraph = tf.paragraphs[0]
    paragraph.alignment = align
    paragraph.font.name = font_name
    paragraph.font.size = Pt(font_size)
    paragraph.font.color.rgb = color
    paragraph.font.bold = bold
    return box


def add_title_slide(prs: Presentation, slide_content: SlideContent) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_background(slide)

    add_textbox(
        slide,
        0.75,
        0.65,
        4.5,
        0.4,
        "DATA MINING RESEARCH 2026",
        font_name="Aptos",
        font_size=12,
        color=TEXT_MUTED,
        bold=True,
    )

    add_textbox(
        slide,
        0.75,
        1.35,
        7.9,
        1.4,
        slide_content.title,
        font_name="Aptos Display",
        font_size=28,
        color=TEXT_MAIN,
        bold=True,
    )

    add_textbox(
        slide,
        0.75,
        2.85,
        8.4,
        0.9,
        slide_content.subtitle,
        font_name="Aptos",
        font_size=20,
        color=TEXT_BODY,
    )

    meta_text = " ".join(slide_content.extras).strip()
    if meta_text:
        add_textbox(
            slide,
            0.75,
            3.85,
            7.2,
            0.7,
            meta_text,
            font_name="Aptos",
            font_size=16,
            color=ACCENT,
            bold=True,
        )

    panel = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(8.95), Inches(1.2), Inches(3.35), Inches(4.3)
    )
    panel.fill.solid()
    panel.fill.fore_color.rgb = PANEL_FILL
    panel.line.color.rgb = ACCENT_SOFT

    add_textbox(
        slide,
        9.25,
        1.65,
        2.75,
        3.3,
        "ID3\n|- Gain\n|  |- Leaf\n|  |- Leaf\n|- Tree\n   |- Leaf\n   |- Leaf",
        font_name="Consolas",
        font_size=20,
        color=TEXT_MUTED,
    )


def add_content_slide(prs: Presentation, slide_content: SlideContent, slide_number: int, total_slides: int) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_background(slide)

    add_textbox(
        slide,
        0.75,
        0.65,
        2.0,
        0.35,
        f"SLIDE {slide_number:02d}",
        font_name="Aptos",
        font_size=12,
        color=TEXT_MUTED,
        bold=True,
    )

    add_textbox(
        slide,
        0.75,
        1.1,
        11.6,
        0.8,
        slide_content.title,
        font_name="Aptos Display",
        font_size=24,
        color=TEXT_MAIN,
        bold=True,
    )

    if slide_content.subtitle:
        add_textbox(
            slide,
            0.75,
            1.85,
            11.4,
            0.55,
            slide_content.subtitle,
            font_name="Aptos",
            font_size=18,
            color=TEXT_MUTED,
        )

    panel = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(0.75), Inches(2.58), Inches(11.85), Inches(4.0)
    )
    panel.fill.solid()
    panel.fill.fore_color.rgb = PANEL_FILL
    panel.line.color.rgb = ACCENT_SOFT

    body_box = slide.shapes.add_textbox(Inches(1.0), Inches(2.88), Inches(11.2), Inches(3.35))
    tf = body_box.text_frame
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE

    bullet_lines = slide_content.bullets[:10]
    if not bullet_lines and slide_content.extras:
        bullet_lines = [(0, extra) for extra in slide_content.extras]

    for index, (level, text) in enumerate(bullet_lines):
        prefix = "- " if level == 0 else "  - "
        if index == 0:
            paragraph = tf.paragraphs[0]
        else:
            paragraph = tf.add_paragraph()

        paragraph.text = f"{prefix}{text}"
        paragraph.font.name = "Aptos"
        paragraph.font.size = Pt(18 if level == 0 else 16)
        paragraph.font.color.rgb = TEXT_BODY
        paragraph.space_after = Pt(8)

    add_textbox(
        slide,
        11.35,
        6.8,
        1.05,
        0.3,
        f"{slide_number:02d}/{total_slides:02d}",
        font_name="Consolas",
        font_size=11,
        color=TEXT_MUTED,
        align=PP_ALIGN.RIGHT,
    )


def build_presentation(slides: list[SlideContent], output_path: Path) -> None:
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    add_title_slide(prs, slides[0])

    total_slides = len(slides)
    for index, slide_content in enumerate(slides[1:], start=2):
        add_content_slide(prs, slide_content, index, total_slides)

    prs.save(output_path)


def main() -> None:
    slides = parse_presentation(SOURCE_MD)
    if not slides:
        raise SystemExit("No slides found in presentation.md")

    build_presentation(slides, OUTPUT_PPTX)
    print(f"Generated {OUTPUT_PPTX}")


if __name__ == "__main__":
    main()
